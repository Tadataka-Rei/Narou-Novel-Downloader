from __future__ import annotations

import os
from docx import Document


def sanitize_filename(name: str) -> str:
    # remove invalid filename characters on Windows because windows suck
    return "".join(c for c in name if c not in r'<>:"/\\|?*')


def save_chapter(
    output_folder: str,
    chapter_num: int,
    title: str,
    content: str,
    file_format: str = "txt",
) -> str:
    # Normalize extension format
    ext = file_format.lower().lstrip(".")
    if not ext:
        ext = "txt"

    filename = f"{chapter_num:04d}_{sanitize_filename(title)}.{ext}"
    path = os.path.join(output_folder, filename)

    if ext in ("docx", "doc"):
        doc = Document()
        doc.add_heading(title, level=1)
        for paragraph in content.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph)
        doc.save(path)
    else:
        with open(path, "w", encoding="utf-8") as f:
            f.write(title + "\n\n")
            f.write(content)

    return path